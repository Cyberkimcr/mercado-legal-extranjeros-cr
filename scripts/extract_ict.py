"""Extrae y valida los datos del Instituto Costarricense de Turismo (ICT).

Fuentes (data/raw/ict):
  - Series_sitio_web_ICT_2025.xlsx ............ llegadas anuales por país 2017-2025 y totales históricos
  - 2025-12_Llegadas_internacionales.pdf ...... llegadas mensuales 2025 por país y puerto
  - 2026-08_Llegadas_internacionales.pdf ...... llegadas mensuales ene-ago 2026 por país y puerto
  - NR_Daniel_Oduber_2025.pdf ................. encuesta a no residentes que salen por Liberia (LIR)
  - powerbi_perfil_turista_vacaciones_sjo.json  datos extraídos del tablero Power BI del ICT
                                               "Turistas por vacaciones" (SJO, 2011-2025)

Cada tabla se valida contra los totales publicados; el script se detiene si algo no cuadra.
"""
import csv, json, re, collections
from pathlib import Path
import openpyxl, pdfplumber

ROOT = Path(__file__).resolve().parents[1]
RAW, OUT = ROOT / "data/raw/ict", ROOT / "data/clean"
MESES = ['Enero','Febrero','Marzo','Abril','Mayo','Junio','Julio','Agosto','Setiembre','Octubre','Noviembre','Diciembre']
ZONAS_PDF = {'AMÉRICA CENTRAL','AMÉRICA DEL NORTE','AMÉRICA DEL SUR','ASIA Y EL PACÍFICO','CARIBE','EUROPA','ORIENTE MEDIO','ÁFRICA'}


def write(name, header, rows):
    with open(OUT / name, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
    print(f'  -> {name}: {len(rows)} filas')


def check(cond, msg):
    if not cond:
        raise SystemExit(f'VALIDACIÓN FALLIDA: {msg}')


# 1. Llegadas anuales por país (Excel) -------------------------------------------------
def anuales():
    wb = openpyxl.load_workbook(RAW / 'Series_sitio_web_ICT_2025.xlsx', data_only=True, read_only=True)
    zonas = {'AMÉRICA DEL NORTE','AMÉRICA CENTRAL','AMÉRICA DEL SUR','CARIBE','EUROPA'}
    rows = []
    for sh, via in [('Todas_vías', 'Todas las vías'), ('Vía_aérea', 'Vía aérea')]:
        R = [list(r) for r in wb[sh].iter_rows(values_only=True)]
        hi = next(i for i, r in enumerate(R) if 2017 in r)
        yc = [(j, v) for j, v in enumerate(R[hi]) if isinstance(v, int) and 2000 < v < 2100]
        lc = min(j for j, _ in yc) - 1
        zona = None
        for r in R[hi + 1:]:
            lab = next((r[j] for j in range(lc + 1) if isinstance(r[j], str) and r[j].strip()), None)
            if lab is None:
                continue
            lab = lab.strip()
            if lab.startswith(('1/', '2/', 'Fuente')):
                break
            vals = [r[j] if j < len(r) else None for j, _ in yc]
            if not any(isinstance(v, (int, float)) for v in vals):
                continue
            name = re.sub(r'\d$', '', lab)  # quita llamadas de nota (Puerto Rico1, OTRAS ZONAS2)
            if name == 'OTRAS ZONAS':  # el Excel no desglosa Asia, África ni Oriente Medio
                name, zona = 'Otras zonas (Asia, África y Oriente Medio)', 'Otras Zonas'
            if lab == 'TOTAL':
                nivel = 'total'
            elif lab in zonas:
                nivel, zona = 'zona', name.title()
            else:
                nivel = 'pais'
            for (_, y), v in zip(yc, vals):
                if isinstance(v, (int, float)) and v != '-':
                    rows.append([y, via, nivel, zona if nivel == 'pais' else name.title(), name, int(v)])
    # validación: países (incluida la agrupación "Otras zonas") = TOTAL
    s, T = collections.Counter(), {}
    for y, via, nivel, zona, name, v in rows:
        if nivel == 'pais':
            s[(y, via)] += v
        if nivel == 'total':
            T[(y, via)] = v
    check(len(T) == 18 and all(s[k] == T[k] for k in T), 'llegadas anuales: países no suman el TOTAL')
    write('ict_llegadas_anuales.csv', ['anio','via','nivel','zona','pais_o_zona','llegadas'], rows)

    hist = []
    for sh, via in [('TV_1951-2025', 'Todas las vías'), ('VA_1976-2025', 'Vía aérea')]:
        for r in wb[sh].iter_rows(values_only=True):
            v = [x for x in r if x is not None]
            if len(v) >= 2 and isinstance(v[0], int) and 1900 < v[0] < 2100 and isinstance(v[1], (int, float)):
                hist.append([v[0], via, int(v[1])])
    write('ict_llegadas_historicas_totales.csv', ['anio','via','llegadas'], hist)


# 2. Llegadas mensuales por país y puerto (PDF) ---------------------------------------
def mensuales():
    cuadros = {1: 'Todas las vías', 3: 'Vía aérea', 5: 'Aeropuerto Juan Santamaría (SJO)', 7: 'Aeropuerto Daniel Oduber (LIR)'}
    rows = []
    for fname, anio in [('2025-12_Llegadas_internacionales.pdf', 2025), ('2026-08_Llegadas_internacionales.pdf', 2026)]:
        with pdfplumber.open(RAW / fname) as pdf:
            for cu, puerto in cuadros.items():
                ws = pdf.pages[cu].extract_words()  # el cuadro N está en la página N+1
                hdr = {}
                for w in ws:
                    if w['text'] in MESES and w['text'] not in hdr:
                        hdr[w['text']] = (w['x0'] + w['x1']) / 2
                    if w['text'].startswith('Ene-'):
                        hdr['TOTAL'] = (w['x0'] + w['x1']) / 2
                htop = max(w['top'] for w in ws if w['text'] in MESES)
                cols = sorted(hdr.items(), key=lambda kv: kv[1])
                lines = {}
                for w in ws:
                    if w['top'] <= htop + 3:
                        continue
                    k = round(w['top'])
                    key = next((t for t in lines if abs(t - k) <= 2), k)
                    lines.setdefault(key, []).append(w)
                zona = None
                for t in sorted(lines):
                    L = sorted(lines[t], key=lambda w: w['x0'])
                    lab = [w for w in L if not re.fullmatch(r'[\d.,-]+', w['text']) or w['x1'] < 180]
                    nums = [w for w in L if w not in lab]
                    label = ' '.join(w['text'] for w in lab).strip()
                    if not label or label.startswith(('Fuente', '1/', '2/', 'Nota', '*')):
                        continue
                    # el ICT separa miles con espacio: se unen los grupos de 3 dígitos contiguos
                    groups, cur = [], None
                    for w in nums:
                        if cur and w['x0'] - cur['x1'] <= 3.5 and re.fullmatch(r'\d{3}', w['text']):
                            cur['text'] += w['text']; cur['x1'] = w['x1']
                        else:
                            cur = {'text': w['text'], 'x0': w['x0'], 'x1': w['x1']}; groups.append(cur)
                    vals = {}
                    for g in groups:
                        c = (g['x0'] + g['x1']) / 2
                        col = min(cols, key=lambda kv: abs(kv[1] - c))[0]
                        check(col not in vals, f'{fname} cuadro {cu}: dos valores en la misma columna ({label})')
                        vals[col] = int(g['text'])
                    if not vals:
                        continue
                    if label == 'TOTAL':
                        nivel = 'total'
                    elif label in ZONAS_PDF:
                        nivel, zona = 'zona', label.title()
                    else:
                        nivel = 'pais'
                    mes = [vals.get(m) for m in MESES]
                    if 'TOTAL' in vals:
                        check(sum(v for v in mes if v) == vals['TOTAL'], f'{fname} cuadro {cu}: meses no suman el total de {label}')
                    name = label.title() if (nivel == 'pais' and label.isupper()) else label
                    for i, v in enumerate(mes):
                        if v is not None:
                            rows.append([anio, i + 1, puerto, nivel, zona if nivel == 'pais' else label.title(), name, v])
    s, T, zs, Z = collections.Counter(), {}, collections.Counter(), {}
    for anio, mes, puerto, nivel, zona, name, v in rows:
        if nivel == 'pais':
            s[(anio, mes, puerto)] += v; zs[(anio, mes, puerto, zona)] += v
        elif nivel == 'total':
            T[(anio, mes, puerto)] = v
        else:
            Z[(anio, mes, puerto, zona)] = v
    check(len(T) == 80 and all(s[k] == T[k] for k in T), 'llegadas mensuales: países no suman el TOTAL')
    check(all(zs[k] == Z[k] for k in Z), 'llegadas mensuales: países no suman su zona')
    write('ict_llegadas_mensuales.csv', ['anio','mes','puerto','nivel','zona','pais_o_zona','llegadas'], rows)


# 3. Encuesta a no residentes, aeropuerto Daniel Oduber (LIR) 2025 ---------------------
def encuesta_lir():
    txt = ''
    with pdfplumber.open(RAW / 'NR_Daniel_Oduber_2025.pdf') as p:
        for pg in p.pages[1:3]:
            txt += pg.extract_text() + '\n'
    txt = txt.replace('Publicación sobre Costa Rica (artículo, anuncio, reportaje o\n18,46',
                      'Publicación sobre Costa Rica (artículo, anuncio, reportaje o comentarios en redes sociales) 18,46')
    txt = txt.replace('\ncomentarios en redes sociales)\n', '\n').replace('área\n100,0', 'área protegida? 100,0').replace('\nprotegida?\n', '\n')
    num = r'(\d[\d ]*,\d+)'
    temas = {'Motivo principal de la visita (%)': ('motivo_visita', '%', False),
             'Modalidad del viaje (%)': ('modalidad_viaje', '%', False),
             '¿Es ésta su primera visita a Costa Rica?': ('primera_visita', '%', False),
             'Manera en que se interesó (%) [Respuesta Múltiple]': ('como_se_intereso', '%', True),
             '¿Quién le acompañó en el viaje? (%)': ('acompanante', '%', False),
             'Establecimiento de Hospedaje (%) [Respuesta Múltiple]': ('hospedaje', '%', True),
             '¿Visitó algún parque nacional, reserva biológica o área protegida?': ('visito_area_protegida', '%', False),
             'Actividades realizadas - grupos (%) [Respuesta Múltiple]': ('grupo_actividad', '%', True)}
    regiones = ['Total', 'Estados Unidos', 'Canadá', 'Europa']
    rows, tema = [], None
    for line in txt.split('\n'):
        line = line.strip()
        head = re.sub(r'\s+100,0.*$', '', line)
        if head in temas or line in temas:
            tema = temas.get(head) or temas[line]; continue
        m = re.match(r'^(.*?)\s+' + r'\s+'.join([num] * 4) + '$', line)
        if not m:
            continue
        cat = m.group(1)
        vals = [float(m.group(i).replace(' ', '').replace(',', '.')) for i in range(2, 6)]
        t = {'Estadía media (En noches)': ('estadia_media', 'noches', False),
             'Gasto medio por persona (En $EEUU)': ('gasto_medio_viaje', 'USD por persona', False)}.get(cat)
        if cat.startswith('Promedio de veces'):
            t = ('veces_visita_previa', 'promedio', False)
        t = t or tema
        nivel = 'subtotal' if t[0] == 'motivo_visita' and cat in ('Motivos personales', 'Negocios y motivos profesionales') else 'detalle'
        for r, v in zip(regiones, vals):
            rows.append(['Daniel Oduber (LIR)', 2025, r, t[0], cat, nivel, v, t[1], t[2]])
    for t in ['modalidad_viaje', 'primera_visita', 'acompanante', 'visito_area_protegida']:
        check(abs(sum(r[6] for r in rows if r[3] == t and r[2] == 'Total') - 100) < 0.05, f'encuesta LIR: {t} no suma 100%')
    check(len(rows) == 220, 'encuesta LIR: se esperaban 220 datos')
    write('ict_encuesta_no_residentes_lir.csv',
          ['aeropuerto','anio','region_residencia','tema','categoria','nivel','valor','unidad','respuesta_multiple'], rows)


# 4. Perfil del turista de vacaciones (tablero Power BI, SJO) --------------------------
def perfil_sjo():
    d = json.load(open(RAW / 'powerbi_perfil_turista_vacaciones_sjo.json', encoding='utf-8'))
    P = ['2011-2013', '2014-2016', '2017-2019', '2023', '2024', '2025']
    TEMA = {"Promedio de revisita por año": ("revisita", "% que ya había visitado CR"), "Escolaridad": ("escolaridad", "%"),
            "Estado civil": ("estado_civil", "%"), "Edad promedio por año": ("edad_promedio", "años"), "Edades": ("grupo_edad", "%"),
            "Promedio de estadía (noches)": ("estadia_promedio", "noches"), "Gasto medio diario por persona": ("gasto_diario", "USD por persona por día"),
            "Modalidad del viaje (%)": ("modalidad_viaje", "%"), "Actividad": ("actividad", "% que la realizó"),
            "Acompañante": ("acompanante", "%"), "Manera en que se interesó (%)": ("como_se_intereso", "%")}
    rows = []
    for r in d['rows']:
        m, t, c, v = r.split('|')
        tema, unidad = TEMA[d['temas'][int(t)]]
        for p, x in zip(P, v.split(';')):
            if x:
                rows.append([d['mk'][int(m)], tema, d['cats'][int(c)] or tema, p, float(x), unidad, ''])
    # el tablero repite la columna 2023 en "Actividad" de Colombia; la segunda copia se recuperó aparte
    col23 = {'Sol y playa': 64.70, 'Disfrute de la gastronomía local': 64.70, 'Observación de la flora y la fauna': 29.40,
             'Visita a los volcanes': 23.50, 'Compras (artesanías, arte, licores, café)': 23.50, 'Trekking - Hiking (terreno más dificil)': 26.50,
             'Aguas termales': 26.50, 'Visita a museos, teatros, galerías, arte': 14.70, 'Visita a comunidades o iglesias': 11.80, 'Surf': 11.80,
             'Snorkel': 5.90, 'Puentes colgantes': 5.90, 'Paseos en Bote, Moto acúatica, Sky acuático': 8.80,
             'Bienestar (tratamientos/masajes terapéuticos)': 0.00}
    rows += [['Colombia', 'actividad', k, '2023', v, '% que la realizó', ''] for k, v in col23.items()]
    for r in rows:
        if r[0] == 'Colombia' and r[1] == 'actividad' and r[3] == '2023':
            r[6] = 'La fuente duplica la columna 2023; se unieron ambas'
        if r[0] == 'Suiza' and r[2] == 'Soltero' and r[3] == '2025':
            r[6] = 'Valor tal como aparece en la fuente (730); probable error de digitación'
    check(len(rows) == 3436 and abs(sum(r[4] for r in rows) - 92849.61) < 0.01, 'perfil SJO: no coincide con el tablero (3436 datos, suma 92849.61)')
    rows.sort(key=lambda r: (r[0], r[1], r[2], P.index(r[3])))
    write('ict_perfil_turista_vacaciones_sjo.csv', ['mercado','tema','categoria','periodo','valor','unidad','nota'], rows)


if __name__ == '__main__':
    print('ICT:')
    anuales(); mensuales(); encuesta_lir(); perfil_sjo()
